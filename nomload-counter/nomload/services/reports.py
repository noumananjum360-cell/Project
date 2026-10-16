"""PDF reports (ReportLab), Excel export (openpyxl) and dashboard statistics."""
from __future__ import annotations

import datetime as dt
import json
import time
from pathlib import Path

from nomload import APP_NAME, __version__
from nomload.config import assets_dir
from nomload.data.db import Database

NAVY = "#1B2A4E"
AMBER = "#E8A33D"
GREEN = "#1F9D55"
RED = "#D64545"
GREY = "#5A6475"
STATUS_LABEL = {"cleared": "CLEARED", "mismatch": "MISMATCH", "released": "RELEASED (supervisor)",
                "active": "LOADING", "cancelled": "CANCELLED"}
STATUS_COLOR = {"cleared": GREEN, "mismatch": RED, "released": AMBER, "active": NAVY, "cancelled": GREY}


def day_bounds(day: dt.date) -> tuple[float, float]:
    start = time.mktime(day.timetuple())
    return start, start + 86400


def fmt_ts(ts: float | None) -> str:
    return time.strftime("%d-%b-%Y %H:%M:%S", time.localtime(ts)) if ts else "-"


def fmt_dur(s: float) -> str:
    s = int(s)
    return f"{s // 3600}h {s % 3600 // 60:02d}m" if s >= 3600 else f"{s // 60}m {s % 60:02d}s"


# ===================================================================== statistics

def _without_caption(paths: list[str]) -> list[str]:
    """Copies of the photos without the caption strip burned in at the bottom (older photos
    carried technical details there). The originals stay untouched as evidence."""
    import tempfile

    import cv2
    out = []
    tmp = Path(tempfile.mkdtemp(prefix="nomload_cert_"))
    for i, p in enumerate(paths):
        img = cv2.imread(p)
        if img is None or img.shape[0] < 60:
            out.append(p)
            continue
        dst = tmp / f"{i:02d}.jpg"
        cv2.imwrite(str(dst), img[: img.shape[0] - 22], [cv2.IMWRITE_JPEG_QUALITY, 90])
        out.append(str(dst))
    return out

class Stats:
    def __init__(self, db: Database):
        self.db = db

    def today(self) -> dict:
        t0, t1 = day_bounds(dt.date.today())
        done = ("cleared", "mismatch", "released")
        q = self.db.q
        trucks = self.db.scalar(f"SELECT COUNT(*) FROM sessions WHERE deleted=0 AND started_at>=? AND started_at<? "
                                f"AND status IN {done}", (t0, t1)) or 0
        active = self.db.scalar("SELECT COUNT(*) FROM sessions WHERE deleted=0 AND status='active'") or 0
        mism = self.db.scalar("SELECT COUNT(*) FROM sessions WHERE deleted=0 AND started_at>=? AND started_at<? "
                              "AND status IN ('mismatch','released')", (t0, t1)) or 0
        units = q("SELECT p.name, SUM(e.direction*e.units) n FROM count_events e JOIN sessions s ON s.id=e.session_id "
                  "JOIN products p ON p.id=e.product_id WHERE s.deleted=0 AND s.started_at>=? AND s.started_at<? "
                  "AND e.review IN ('auto','accepted','manual') GROUP BY p.name ORDER BY n DESC", (t0, t1))
        total = sum(r["n"] or 0 for r in units)
        avg = self.db.scalar("SELECT AVG(ended_at-started_at) FROM sessions WHERE deleted=0 AND ended_at IS NOT NULL "
                             "AND started_at>=? AND started_at<?", (t0, t1))
        pending = self.db.scalar("SELECT COUNT(*) FROM count_events e JOIN sessions s ON s.id=e.session_id "
                                 "WHERE e.review='pending' AND s.deleted=0") or 0
        return {"trucks": trucks, "active": active, "mismatches": mism, "units_total": total,
                "units_by_product": [(r["name"], r["n"] or 0) for r in units], "avg_load_s": avg or 0,
                "pending_reviews": pending}

    def avg_time_per_bay(self, days: int = 7) -> list[tuple[str, float]]:
        since = time.time() - days * 86400
        return [(r["name"], (r["a"] or 0) / 60.0) for r in self.db.q(
            "SELECT b.name, AVG(s.ended_at-s.started_at) a FROM sessions s JOIN bays b ON b.id=s.bay_id "
            "WHERE s.deleted=0 AND s.ended_at IS NOT NULL AND s.started_at>=? GROUP BY b.name ORDER BY b.name", (since,))]

    def busiest_hours(self, days: int = 7) -> list[int]:
        since = time.time() - days * 86400
        hours = [0] * 24
        for r in self.db.q("SELECT started_at FROM sessions WHERE deleted=0 AND started_at>=?", (since,)):
            hours[time.localtime(r["started_at"]).tm_hour] += 1
        return hours

    def trucks_per_day(self, days: int = 14) -> list[tuple[str, int, int]]:
        out = []
        today = dt.date.today()
        for i in range(days - 1, -1, -1):
            d = today - dt.timedelta(days=i)
            t0, t1 = day_bounds(d)
            ok = self.db.scalar("SELECT COUNT(*) FROM sessions WHERE deleted=0 AND started_at>=? AND started_at<? "
                                "AND status='cleared'", (t0, t1)) or 0
            bad = self.db.scalar("SELECT COUNT(*) FROM sessions WHERE deleted=0 AND started_at>=? AND started_at<? "
                                 "AND status IN ('mismatch','released')", (t0, t1)) or 0
            out.append((d.strftime("%m-%d"), ok, bad))
        return out

    def insights(self) -> list[str]:
        out = []
        t = self.today()
        hrs = self.busiest_hours()
        if any(hrs):
            h = max(range(24), key=lambda i: hrs[i])
            out.append(f"Busiest loading hour over the last 7 days: {h:02d}:00-{(h + 1) % 24:02d}:00 ({hrs[h]} trucks).")
        bays = self.avg_time_per_bay()
        if bays:
            slow = max(bays, key=lambda b: b[1])
            out.append(f"Slowest bay on average: {slow[0]} ({slow[1]:.0f} min per truck).")
        week = self.db.q("SELECT status, COUNT(*) n FROM sessions WHERE deleted=0 AND started_at>=? GROUP BY status",
                         (time.time() - 7 * 86400,))
        tot = sum(r["n"] for r in week if r["status"] != "active")
        bad = sum(r["n"] for r in week if r["status"] in ("mismatch", "released"))
        if tot:
            out.append(f"Match rate over 7 days: {100 * (tot - bad) / tot:.0f}% ({bad} of {tot} trucks had a mismatch).")
        if t["pending_reviews"]:
            out.append(f"{t['pending_reviews']} low-confidence count(s) are waiting for supervisor review.")
        if not out:
            out.append("No loading activity recorded yet. Start a session from Live Bays.")
        return out


# ===================================================================== PDF
def _styles():
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    ss = getSampleStyleSheet()
    ss.add(ParagraphStyle("H1x", parent=ss["Heading1"], fontName="Helvetica-Bold", fontSize=17,
                          textColor=NAVY, spaceAfter=4))
    ss.add(ParagraphStyle("Small", parent=ss["Normal"], fontSize=8, textColor=GREY))
    ss.add(ParagraphStyle("Cell", parent=ss["Normal"], fontSize=9, leading=11))
    ss.add(ParagraphStyle("Banner", parent=ss["Normal"], fontName="Helvetica-Bold", fontSize=16,
                          textColor="#FFFFFF", alignment=1))
    return ss


class _Doc:
    def __init__(self, path, title: str, company: dict, watermark: bool):
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.units import mm
        from reportlab.platypus import SimpleDocTemplate
        self.path = str(path)
        self.title = title
        self.company = company
        self.watermark = watermark
        self.mm = mm
        self.doc = SimpleDocTemplate(self.path, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm,
                                     topMargin=34 * mm, bottomMargin=18 * mm, title=title,
                                     author=company.get("name", ""), creator=f"{APP_NAME} {__version__}")

    def _page(self, c, doc):
        from reportlab.lib.colors import HexColor
        from reportlab.lib.pagesizes import A4
        mm = self.mm
        w, h = A4
        c.saveState()
        logo = self.company.get("logo")
        x = 16 * mm
        has_logo = False
        if logo and Path(logo).exists():
            try:
                from reportlab.lib.utils import ImageReader
                iw, ih = ImageReader(logo).getSize()
                lh = 16 * mm                                  # logo height
                lw = min(70 * mm, lh * iw / max(1, ih))       # keep its shape, at most 70 mm wide
                lh = lw * ih / max(1, iw)
                c.drawImage(logo, x, h - 11 * mm - lh, width=lw, height=lh, mask="auto")
                x += lw + 4 * mm
                has_logo = True
            except Exception:
                pass
        show_name = self.company.get("show_name") or not has_logo
        if show_name and self.company.get("name"):
            c.setFillColor(HexColor(NAVY))
            c.setFont("Helvetica-Bold", 13)
            c.drawString(x, h - 17 * mm, self.company.get("name", ""))
            addr_y = h - 22 * mm
        else:
            addr_y = h - 19.5 * mm
        c.setFont("Helvetica", 8.5)
        c.setFillColor(HexColor(GREY))
        addr, room = self.company.get("address", ""), (w - 64 * mm) - x   # stop before the logo on the right
        while addr and c.stringWidth(addr, "Helvetica", 8.5) > room:
            addr = addr[:-1]
        c.drawString(x, addr_y, addr)
        wm = assets_dir() / "wordmark.png"
        if wm.exists():
            c.drawImage(str(wm), w - 62 * mm, h - 26 * mm, width=46 * mm, height=11.5 * mm, mask="auto")
        c.setStrokeColor(HexColor(AMBER))
        c.setLineWidth(2)
        c.line(16 * mm, h - 30 * mm, w - 16 * mm, h - 30 * mm)
        c.setFont("Helvetica", 7.5)
        c.setFillColor(HexColor(GREY))
        c.drawString(16 * mm, 10 * mm, f"{self.title}  |  {fmt_ts(time.time())}")
        c.drawRightString(w - 16 * mm, 10 * mm, f"Page {doc.page}")
        if self.watermark:
            c.setFont("Helvetica-Bold", 90)
            c.setFillColor(HexColor("#D64545"))
            c.setFillAlpha(0.12)
            c.translate(w / 2, h / 2)
            c.rotate(40)
            c.drawCentredString(0, 0, "TRIAL")
        c.restoreState()

    def build(self, story):
        self.doc.build(story, onFirstPage=self._page, onLaterPages=self._page)
        return Path(self.path)


def _table(rows, widths, header=True, zebra=True, align_right_from: int | None = None):
    from reportlab.lib.colors import HexColor
    from reportlab.platypus import Table, TableStyle
    t = Table(rows, colWidths=widths, repeatRows=1 if header else 0)
    st = [("FONT", (0, 0), (-1, -1), "Helvetica", 9), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
          ("BOTTOMPADDING", (0, 0), (-1, -1), 5), ("TOPPADDING", (0, 0), (-1, -1), 5),
          ("LINEBELOW", (0, 0), (-1, -1), 0.25, HexColor("#E3E6EC"))]
    if header:
        st += [("BACKGROUND", (0, 0), (-1, 0), HexColor(NAVY)), ("TEXTCOLOR", (0, 0), (-1, 0), HexColor("#FFFFFF")),
               ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 9)]
    if zebra:
        for i in range(1 if header else 0, len(rows)):
            if i % 2 == 0:
                st.append(("BACKGROUND", (0, i), (-1, i), HexColor("#F5F7FB")))
    if align_right_from is not None:
        st.append(("ALIGN", (align_right_from, 0), (-1, -1), "RIGHT"))
    t.setStyle(TableStyle(st))
    return t


def _kv_table(pairs, widths):
    from reportlab.lib.colors import HexColor
    from reportlab.platypus import Table, TableStyle
    t = Table(pairs, colWidths=widths)
    t.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", 9.5), ("FONT", (0, 0), (0, -1), "Helvetica-Bold", 9.5),
                           ("TEXTCOLOR", (0, 0), (0, -1), HexColor(GREY)), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                           ("BOTTOMPADDING", (0, 0), (-1, -1), 3), ("TOPPADDING", (0, 0), (-1, -1), 3)]))
    return t


def _banner(text, color, ss):
    from reportlab.lib.colors import HexColor
    from reportlab.platypus import Paragraph, Table, TableStyle
    t = Table([[Paragraph(text, ss["Banner"])]], colWidths=[178 * 2.8346])
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), HexColor(color)), ("TOPPADDING", (0, 0), (-1, -1), 8),
                           ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
    return t


class ReportService:
    def __init__(self, db: Database, settings, sessions, license_status_fn=None):
        self.db = db
        self.settings = settings
        self.sessions = sessions
        self.license_status_fn = license_status_fn

    def _company(self) -> dict:
        return {"name": self.settings.get("company_name") or "", "address": self.settings.get("company_address") or "",
                "logo": self.settings.get("company_logo") or "",
                "show_name": bool(self.settings.get("report_show_company_name"))}

    def _wm(self) -> bool:
        try:
            return bool(self.license_status_fn and self.license_status_fn().watermark)
        except Exception:
            return True

    # ------------------------------------------------------------------ certificate
    def loading_certificate(self, sid: int, out: str | Path) -> Path:
        from reportlab.platypus import Image, Paragraph, Spacer
        ss = _styles()
        s = self.sessions.summary(sid)
        mm = 2.8346
        story = [Paragraph("Truck Loading Certificate", ss["H1x"]),
                 Paragraph(f"Session #{s.id} &nbsp;&middot;&nbsp; {s.bay_name}", ss["Small"]), Spacer(1, 6),
                 _banner(STATUS_LABEL.get(s.status, s.status.upper()), STATUS_COLOR.get(s.status, NAVY), ss), Spacer(1, 8)]
        gp = s.gate_pass
        pairs = [["Gate pass", gp.number if gp else "- (no gate pass)"],
                 ["Customer / distributor", gp.customer if gp and gp.customer else "-"],
                 ["Destination", gp.destination if gp and gp.destination else "-"],
                 ["Vehicle (gate pass)", gp.vehicle_no if gp and gp.vehicle_no else "-"],
                 ["Vehicle number", s.plate_text or "-"],
                 ["Loading started", fmt_ts(s.started_at)], ["Loading finished", fmt_ts(s.ended_at)],
                 ["Loading time", fmt_dur(s.duration_s)], ["Started by", s.started_by or "-"],
                 ["Closed by", s.closed_by or "-"]]
        if s.status == "released":
            pairs += [["Released by supervisor", s.released_by], ["Release reason", Paragraph(s.release_reason, ss["Cell"])]]
        kv = _kv_table(pairs, [48 * mm, 80 * mm])
        if s.plate_image and Path(s.plate_image).exists():
            from reportlab.platypus import Table
            img = Image(s.plate_image, width=46 * mm, height=16 * mm, kind="proportional")
            story.append(Table([[kv, img]], colWidths=[130 * mm, 48 * mm], style=[("VALIGN", (0, 0), (-1, -1), "TOP")]))
        else:
            story.append(kv)
        story += [Spacer(1, 10), Paragraph("<b>Quantities</b>", ss["Normal"]), Spacer(1, 4)]
        rows = [["Product", "Expected", "Loaded", "Difference", "Result"]]
        for l in s.lines:
            rows.append([Paragraph(l.name, ss["Cell"]), "-" if l.expected is None else str(l.expected), str(l.counted),
                         f"{l.diff:+d}" if l.expected is not None else "-", "OK" if l.ok else "MISMATCH"])
        rows.append(["Total", str(s.total_expected), str(s.total_counted), f"{s.total_counted - s.total_expected:+d}", ""])
        t = _table(rows, [78 * mm, 24 * mm, 24 * mm, 24 * mm, 28 * mm], align_right_from=1)
        story.append(t)
        # snapshots
        snaps = [r["snapshot"] for r in self.db.q("SELECT snapshot FROM count_events WHERE session_id=? AND snapshot "
                                                  "IS NOT NULL ORDER BY ts LIMIT 12", (sid,)) if Path(r["snapshot"]).exists()]
        if snaps:
            from reportlab.platypus import Table
            story += [Spacer(1, 10), Paragraph("<b>Photos</b>", ss["Normal"]),
                      Spacer(1, 4)]
            cells = [Image(p, width=42 * mm, height=26 * mm, kind="proportional") for p in _without_caption(snaps)]
            grid = [cells[i:i + 4] for i in range(0, len(cells), 4)]
            grid[-1] += [""] * (4 - len(grid[-1]))
            story.append(Table(grid, colWidths=[44.5 * mm] * 4))
        story += [Spacer(1, 18), _kv_table([["Gate clerk signature", "____________________"],
                                            ["Supervisor signature", "____________________"],
                                            ["Driver signature", "____________________"]], [48 * mm, 80 * mm])]
        return _Doc(out, "Truck Loading Certificate", self._company(), self._wm()).build(story)

    # ------------------------------------------------------------------ daily summary
    def daily_dispatch(self, day: dt.date, out: str | Path) -> Path:
        from reportlab.platypus import Paragraph, Spacer
        ss = _styles()
        mm = 2.8346
        t0, t1 = day_bounds(day)
        rows_s = self.sessions.search(t0, t1)
        story = [Paragraph(f"Daily Dispatch Summary - {day.strftime('%A %d %B %Y')}", ss["H1x"]), Spacer(1, 6)]
        done = [r for r in rows_s if r["status"] in ("cleared", "mismatch", "released")]
        mism = [r for r in done if r["status"] != "cleared"]
        prod = self.db.q("SELECT p.name, SUM(e.direction*e.units) n FROM count_events e JOIN sessions s ON s.id=e.session_id "
                         "JOIN products p ON p.id=e.product_id WHERE s.deleted=0 AND s.started_at>=? AND s.started_at<? "
                         "AND e.review IN ('auto','accepted','manual') GROUP BY p.name ORDER BY p.name", (t0, t1))
        kpi = [["Trucks loaded", "Units dispatched", "Mismatches", "Cleared first time"],
               [str(len(done)), str(sum(r["n"] or 0 for r in prod)), str(len(mism)),
                f"{(100 * (len(done) - len(mism)) / len(done)):.0f}%" if done else "-"]]
        story += [_table(kpi, [44.5 * mm] * 4, zebra=False), Spacer(1, 10),
                  Paragraph("<b>Units dispatched per product</b>", ss["Normal"]), Spacer(1, 4)]
        story.append(_table([["Product", "Units"]] + [[r["name"], str(r["n"] or 0)] for r in prod] or [["-", "0"]],
                            [130 * mm, 48 * mm], align_right_from=1))
        story += [Spacer(1, 10), Paragraph("<b>Trucks</b>", ss["Normal"]), Spacer(1, 4)]
        rows = [["#", "Time", "Bay", "Gate pass", "Vehicle", "Counted", "Expected", "Status"]]
        for r in sorted(rows_s, key=lambda r: r["started_at"]):
            res = json.loads(r["result"] or "{}")
            rows.append([str(r["id"]), time.strftime("%H:%M", time.localtime(r["started_at"])), r["bay"],
                         r["gate_pass"] or "-", r["plate_text"] or "-", str(res.get("total_counted", "-")),
                         str(res.get("total_expected", "-")), STATUS_LABEL.get(r["status"], r["status"])[:12]])
        story.append(_table(rows, [12 * mm, 14 * mm, 22 * mm, 26 * mm, 26 * mm, 18 * mm, 18 * mm, 42 * mm]))
        return _Doc(out, "Daily Dispatch Summary", self._company(), self._wm()).build(story)

    # ------------------------------------------------------------------ mismatch report
    def mismatch_report(self, d0: dt.date, d1: dt.date, out: str | Path) -> Path:
        from reportlab.platypus import Paragraph, Spacer
        ss = _styles()
        mm = 2.8346
        t0, _ = day_bounds(d0)
        _, t1 = day_bounds(d1)
        rows_s = [r for r in self.sessions.search(t0, t1) if r["status"] in ("mismatch", "released")]
        story = [Paragraph(f"Mismatch Report - {d0.strftime('%d %b %Y')} to {d1.strftime('%d %b %Y')}", ss["H1x"]),
                 Paragraph(f"{len(rows_s)} truck(s) where the counted quantity did not match the gate pass.", ss["Small"]),
                 Spacer(1, 8)]
        rows = [["#", "Date/time", "Bay", "Gate pass", "Vehicle", "Product differences", "Status / reason"]]
        for r in sorted(rows_s, key=lambda r: r["started_at"]):
            s = self.sessions.summary(r["id"])
            diffs = "<br/>".join(f"{l.name}: {l.counted}/{l.expected if l.expected is not None else '-'} ({l.diff:+d})"
                                 for l in s.lines if not l.ok)
            reason = (f"Released by {s.released_by}: {s.release_reason}" if s.status == "released" else "Not released")
            rows.append([str(s.id), fmt_ts(s.started_at)[:17], s.bay_name, s.gate_pass.number if s.gate_pass else "-",
                         s.plate_text or "-", Paragraph(diffs or "-", ss["Cell"]), Paragraph(reason, ss["Cell"])])
        story.append(_table(rows, [10 * mm, 28 * mm, 16 * mm, 20 * mm, 20 * mm, 44 * mm, 40 * mm]))
        return _Doc(out, "Mismatch Report", self._company(), self._wm()).build(story)

    # ------------------------------------------------------------------ excel
    def export_excel(self, d0: dt.date, d1: dt.date, out: str | Path) -> Path:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        t0, _ = day_bounds(d0)
        _, t1 = day_bounds(d1)
        wb = Workbook()
        hdr_font = Font(bold=True, color="FFFFFF")
        hdr_fill = PatternFill("solid", fgColor=NAVY[1:])

        def sheet(ws, header, rows):
            ws.append(header)
            for c in ws[1]:
                c.font, c.fill = hdr_font, hdr_fill
            for r in rows:
                ws.append(list(r))
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = min(45, max(10, max(len(str(c.value or "")) for c in col) + 2))
            ws.freeze_panes = "A2"

        ss = self.sessions.search(t0, t1, include_deleted=False, limit=100000)
        ws = wb.active
        ws.title = "Sessions"
        sheet(ws, ["Session", "Started", "Ended", "Bay", "Gate pass", "Customer", "Vehicle", "Status", "Expected",
                   "Counted", "Pending review"],
              [[r["id"], fmt_ts(r["started_at"]), fmt_ts(r["ended_at"]), r["bay"], r["gate_pass"] or "", r["customer"] or "",
                r["plate_text"], r["status"], json.loads(r["result"] or "{}").get("total_expected"),
                json.loads(r["result"] or "{}").get("total_counted"), json.loads(r["result"] or "{}").get("pending")]
               for r in ss])
        lines = []
        for r in ss:
            for l in self.sessions.lines(r["id"]):
                lines.append([r["id"], r["gate_pass"] or "", l.name, l.expected, l.counted, l.diff, l.allowed, l.pending,
                              "OK" if l.ok else "MISMATCH"])
        sheet(wb.create_sheet("Products"), ["Session", "Gate pass", "Product", "Expected", "Counted", "Difference",
                                            "Tolerance", "Pending", "Result"], lines)
        ids = [r["id"] for r in ss] or [-1]
        ev = self.db.q(f"SELECT e.*, p.name pname FROM count_events e LEFT JOIN products p ON p.id=e.product_id "
                       f"WHERE e.session_id IN ({','.join('?' * len(ids))}) ORDER BY e.ts", ids)
        sheet(wb.create_sheet("Count events"), ["Event", "Session", "Time", "Product", "Class", "Direction", "Units",
                                                "Confidence", "Review", "Reason", "Snapshot"],
              [[e["id"], e["session_id"], fmt_ts(e["ts"]), e["pname"] or "", e["cls"], e["direction"], e["units"],
                round(e["confidence"], 3), e["review"], e["reason"], e["snapshot"] or ""] for e in ev])
        al = self.db.q("SELECT * FROM alerts WHERE ts>=? AND ts<? ORDER BY ts", (t0, t1))
        sheet(wb.create_sheet("Alerts"), ["Time", "Kind", "Severity", "Bay", "Session", "Message"],
              [[fmt_ts(a["ts"]), a["kind"], a["severity"], a["bay_id"], a["session_id"], a["message"]] for a in al])
        wb.save(str(out))
        return Path(out)
