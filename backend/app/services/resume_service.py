"""Resume upload: validation, text extraction, structured extraction and retention."""

from __future__ import annotations

import io
import os
import uuid
import zipfile
from datetime import UTC, datetime
from pathlib import Path

from sqlalchemy.orm import Session

from app.ai import prompts
from app.ai.factory import get_ai_provider
from app.ai.schemas import ResumeExtraction
from app.core.config import get_settings
from app.core.errors import AppError
from app.core.logging import get_logger
from app.models import Resume, User
from app.services.ai_runner import ai_complete
from app.services.skills import extract_skills_from_text

log = get_logger(__name__)

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def detect_type(filename: str, data: bytes) -> str:
    """Validate by extension AND magic bytes; never trust the client content type."""
    ext = Path(filename or "").suffix.lower()
    if ext == ".pdf" and data[:5] == b"%PDF-":
        return PDF
    if ext == ".docx" and data[:4] == b"PK\x03\x04":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                names = set(zf.namelist())
                if "word/document.xml" in names and "[Content_Types].xml" in names:
                    total = sum(i.file_size for i in zf.infolist())
                    if total > 50 * 1024 * 1024:  # zip-bomb guard
                        raise AppError("The document is too large when decompressed.", code="invalid_file")
                    return DOCX
        except zipfile.BadZipFile:
            pass
    raise AppError(
        "Only PDF and DOCX files are supported, and the file content must match its extension.", code="invalid_file_type"
    )


def extract_text(content_type: str, data: bytes) -> str:
    try:
        if content_type == PDF:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise AppError("Password-protected PDFs are not supported.", code="unreadable_file")
            text = "\n".join((page.extract_text() or "") for page in reader.pages[:30])
        else:
            import docx

            document = docx.Document(io.BytesIO(data))
            parts = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            text = "\n".join(parts)
    except AppError:
        raise
    except Exception as exc:  # noqa: BLE001 - parser libraries raise many types
        log.info("resume_parse_failed", error_type=type(exc).__name__)
        raise AppError("The document could not be read. It may be corrupted or malformed.", code="unreadable_file") from exc
    text = text.replace("\x00", "").strip()
    if len(text) < 30:
        raise AppError(
            "Very little text could be extracted. Scanned (image-only) documents are not supported.",
            code="unreadable_file",
        )
    return text[:50000]


def _store_file(user_id: uuid.UUID, data: bytes, ext: str) -> str:
    base = Path(get_settings().UPLOAD_DIRECTORY) / "resumes" / str(user_id)
    base.mkdir(parents=True, exist_ok=True)
    path = base / f"{uuid.uuid4().hex}{ext}"
    with open(path, "wb") as fh:
        fh.write(data)
    os.chmod(path, 0o600)
    return str(path)


def delete_file(resume: Resume) -> None:
    if resume.storage_path:
        try:
            Path(resume.storage_path).unlink(missing_ok=True)
        except OSError:
            log.warning("resume_file_delete_failed")
        resume.storage_path = None
        resume.file_deleted_at = datetime.now(UTC)


def process_upload(db: Session, user: User, filename: str, data: bytes) -> Resume:
    settings = get_settings()
    if len(data) == 0:
        raise AppError("The uploaded file is empty.", code="invalid_file")
    if len(data) > settings.MAX_UPLOAD_SIZE_MB * 1024 * 1024:
        raise AppError(f"File exceeds the {settings.MAX_UPLOAD_SIZE_MB} MB limit.", code="file_too_large", status_code=413)
    ctype = detect_type(filename, data)
    safe_name = Path(filename).name[:255]
    resume = Resume(user_id=user.id, original_filename=safe_name, content_type=ctype, size_bytes=len(data))
    db.add(resume)
    db.flush()
    # Files are processed in memory; they are only written to disk if retention is enabled.
    if settings.RESUME_RETAIN_FILES:
        resume.storage_path = _store_file(user.id, data, ".pdf" if ctype == PDF else ".docx")
    try:
        text = extract_text(ctype, data)
    except AppError as exc:
        resume.status = "failed"
        resume.error = exc.message
        db.commit()
        raise
    resume.extracted_text = text
    keyword_skills = extract_skills_from_text(text)
    try:
        provider = get_ai_provider(db)
        result, _ = ai_complete(
            db,
            provider,
            user_id=user.id,
            operation="resume_extract",
            system=prompts.RESUME_SYSTEM,
            user=prompts.resume_user(text),
            schema=ResumeExtraction,
            temperature=0,
        )
        data_out = result.data.model_dump()
        resume.extraction_method = "ai"
        resume.is_simulated = result.is_simulated
    except AppError as exc:
        # Honest fallback: keyword matching only, clearly labelled.
        log.info("resume_ai_extraction_unavailable", code=exc.code)
        data_out = ResumeExtraction(skills=keyword_skills).model_dump()
        data_out["notice"] = f"AI extraction unavailable ({exc.message}). Only skills were detected by keyword matching."
        resume.extraction_method = "keyword"
    data_out["keyword_skills"] = keyword_skills
    data_out["verification_note"] = "Extracted claims are unverified until you review and confirm them."
    resume.extracted_data = data_out
    resume.status = "parsed"
    db.commit()
    db.refresh(resume)
    return resume
